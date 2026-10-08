# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "pdfplumber==0.11.7",
#   "openpyxl==3.1.5",
# ]
# ///

"""Convert a CIB debit statement to Excel and reconcile a read-only API snapshot."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ZONE = timezone(timedelta(hours=8))
SHEET = "兴业银行流水"
HEADERS = ("顺序号", "交易时间", "记账日期", "摘要", "支/收", "交易金额", "账户余额", "交易地点", "PDF页码", "PDF行号")
PDF_HEADERS = ("交易时间", "记账日期", "摘要", "支/收", "交易金额", "账户余额", "交易地点")
MONEY = re.compile(r"^-?(?:\d{1,3}(?:,\d{3})*|\d+)\.\d{2}$")
INTERNAL_SUMMARIES = {"理财代销", "购汇"}


class ReconcileError(ValueError):
    pass


@dataclass(frozen=True)
class BankRow:
    time: datetime
    posting_date: str
    summary: str
    direction: str
    amount: int
    balance: int
    place: str
    page: int
    row: int


@dataclass(frozen=True)
class LedgerRow:
    time: datetime
    amount: int
    kind: int
    comment: str
    id: str


def cents(value: str, location: str) -> int:
    value = str(value).strip()
    if not MONEY.fullmatch(value):
        raise ReconcileError(f"{location}: invalid amount {value!r}")
    return int(Decimal(value.replace(",", "")) * 100)


def excel_cents(value: object, location: str) -> int:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise ReconcileError(f"{location}: expected a numeric cell")
    scaled = Decimal(str(value)) * 100
    if scaled != scaled.to_integral_value():
        raise ReconcileError(f"{location}: amount has more than two decimals")
    return int(scaled)


def yuan(value: int) -> str:
    sign = "-" if value < 0 else ""
    value = abs(value)
    return f"{sign}{value // 100:,}.{value % 100:02d}"


def clean(value: str | None) -> str:
    return "".join((value or "").split())


def validate(rows: list[BankRow]) -> None:
    if not rows:
        raise ReconcileError("statement has no transactions")
    for previous, current in zip(rows, rows[1:]):
        if current.time < previous.time:
            raise ReconcileError(f"p{current.page}/r{current.row}: transaction order changed")
        if previous.balance + current.amount != current.balance:
            raise ReconcileError(
                f"p{current.page}/r{current.row}: balance should be "
                f"{yuan(previous.balance + current.amount)}, got {yuan(current.balance)}"
            )


def parse_pdf(path: Path) -> list[BankRow]:
    rows: list[BankRow] = []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            # The printed watermark crosses cells at 10/11.2/13.7 pt. The table is 9 pt.
            table_page = page.filter(
                lambda obj: obj.get("object_type") != "char"
                or abs(float(obj.get("size", 0)) - 9) < 0.05
            )
            tables = table_page.extract_tables()
            if len(tables) != 1 or not tables[0]:
                raise ReconcileError(f"page {page_no}: expected one transaction table")
            table = tables[0]
            if tuple(clean(cell) for cell in table[0]) != PDF_HEADERS:
                raise ReconcileError(f"page {page_no}: unexpected table header")
            for row_no, cells in enumerate(table[1:], 1):
                if len(cells) != 7:
                    raise ReconcileError(f"p{page_no}/r{row_no}: expected seven cells")
                location = f"p{page_no}/r{row_no}"
                try:
                    occurred = datetime.strptime(cells[0], "%Y-%m-%d\n%H:%M:%S").replace(tzinfo=ZONE)
                except (TypeError, ValueError) as exc:
                    raise ReconcileError(f"{location}: invalid transaction time") from exc
                posting_date = clean(cells[1])
                if posting_date != occurred.strftime("%Y%m%d"):
                    raise ReconcileError(f"{location}: posting date differs from transaction date")
                direction = clean(cells[3])
                amount = cents(clean(cells[4]), f"{location} amount")
                if (direction == "收" and amount <= 0) or (direction == "支" and amount >= 0):
                    raise ReconcileError(f"{location}: direction and amount disagree")
                if direction not in ("收", "支"):
                    raise ReconcileError(f"{location}: invalid direction")
                rows.append(BankRow(
                    occurred, posting_date, clean(cells[2]), direction, amount,
                    cents(clean(cells[5]), f"{location} balance"), clean(cells[6]), page_no, row_no,
                ))
    rows.reverse()  # CIB prints newest first; balances must be checked oldest first.
    validate(rows)
    return rows


def write_excel(rows: list[BankRow], pdf_path: Path, output: Path) -> None:
    workbook = Workbook()
    info = workbook.active
    info.title = "转换说明"
    info.append(("源文件", pdf_path.name))
    info.append(("SHA-256", hashlib.sha256(pdf_path.read_bytes()).hexdigest()))
    info.append(("时区", "UTC+8"))
    info.append(("规则", "水印按字形大小过滤；原 PDF 倒序展示；金额和余额以分校验"))
    sheet = workbook.create_sheet(SHEET)
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for sequence, row in enumerate(rows, 1):
        sheet.append((sequence, row.time.replace(tzinfo=None), row.posting_date, row.summary,
                      row.direction, row.amount / 100, row.balance / 100, row.place, row.page, row.row))
        sheet.cell(sequence + 1, 2).number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(sequence + 1, 6).number_format = "#,##0.00;[Red]-#,##0.00"
        sheet.cell(sequence + 1, 7).number_format = "#,##0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{len(rows) + 1}"
    for column, width in {"A": 10, "B": 22, "C": 14, "D": 18, "E": 10, "F": 17,
                          "G": 17, "H": 36, "I": 12, "J": 12}.items():
        sheet.column_dimensions[column].width = width
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def read_excel(path: Path) -> list[BankRow]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if SHEET not in workbook:
            raise ReconcileError(f"missing sheet {SHEET}")
        values = workbook[SHEET].iter_rows(values_only=True)
        if tuple(next(values)) != HEADERS:
            raise ReconcileError("Excel headers changed")
        rows: list[BankRow] = []
        for sequence, cells in enumerate(values, 1):
            if len(cells) != len(HEADERS) or cells[0] != sequence:
                raise ReconcileError(f"Excel row {sequence + 1}: sequence or columns changed")
            when = cells[1]
            if not isinstance(when, datetime):
                raise ReconcileError(f"Excel row {sequence + 1}: invalid time")
            rows.append(BankRow(when.replace(tzinfo=ZONE), str(cells[2]), str(cells[3]), str(cells[4]),
                                excel_cents(cells[5], "Excel amount"), excel_cents(cells[6], "Excel balance"),
                                str(cells[7]), int(cells[8]), int(cells[9])))
    finally:
        workbook.close()
    validate(rows)
    return rows


def read_ledger(path: Path, account_id: str) -> list[LedgerRow]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(data, list):
        raise ReconcileError("online snapshot must be a transaction list")
    rows: list[LedgerRow] = []
    for item in data:
        kind = int(item["type"])
        source, destination = str(item.get("sourceAccountId")), str(item.get("destinationAccountId"))
        if kind == 2 and source == account_id:
            amount = int(item["sourceAmount"])
        elif kind == 3 and source == account_id:
            amount = -int(item["sourceAmount"])
        elif kind == 4 and source == account_id:
            amount = -int(item["sourceAmount"])
        elif kind == 4 and destination == account_id:
            amount = int(item["destinationAmount"])
        elif kind == 1 and source == account_id and item.get("balanceDelta") is not None:
            amount = int(item["balanceDelta"])
        else:
            raise ReconcileError(f"unsupported account role or transaction type: {item.get('id')}")
        rows.append(LedgerRow(datetime.fromtimestamp(int(item["time"]), ZONE), amount,
                              kind, str(item.get("comment") or ""), str(item["id"])))
    return rows


def reconcile(bank: list[BankRow], ledger: list[LedgerRow]) -> tuple[list[tuple[BankRow, LedgerRow, str]], list[BankRow], list[LedgerRow]]:
    matched_bank: set[int] = set()
    matched_ledger: set[int] = set()
    pairs: list[tuple[BankRow, LedgerRow, str]] = []
    for phase in ("秒级同额", "同日同额", "七日内同额"):
        candidates = []
        for bi, b in enumerate(bank):
            if bi in matched_bank:
                continue
            for li, l in enumerate(ledger):
                if li in matched_ledger or b.amount != l.amount:
                    continue
                seconds = abs((b.time - l.time).total_seconds())
                days = abs((b.time.date() - l.time.date()).days)
                if (phase == "秒级同额" and seconds == 0 or
                    phase == "同日同额" and days == 0 or
                    phase == "七日内同额" and days <= 7):
                    candidates.append((seconds, bi, li))
        for _, bi, li in sorted(candidates):
            if bi not in matched_bank and li not in matched_ledger:
                matched_bank.add(bi)
                matched_ledger.add(li)
                pairs.append((bank[bi], ledger[li], phase))
    first, last = bank[0].time.date(), bank[-1].time.date()
    return (pairs, [b for i, b in enumerate(bank) if i not in matched_bank],
            sorted((l for i, l in enumerate(ledger)
                    if i not in matched_ledger and first <= l.time.date() <= last), key=lambda l: l.time))


def report(bank: list[BankRow], ledger: list[LedgerRow], account_name: str,
           snapshot: Path, account_balance: int | None = None) -> str:
    excluded = [row for row in bank if row.summary in INTERNAL_SUMMARIES]
    external = [row for row in bank if row.summary not in INTERNAL_SUMMARIES]
    pairs, bank_only, ledger_only = reconcile(external, ledger)
    counts = Counter(phase for _, _, phase in pairs)
    income = sum(row.amount for row in bank if row.amount > 0)
    expense = sum(row.amount for row in bank if row.amount < 0)
    opening = bank[0].balance - bank[0].amount
    bank_net = sum(row.amount for row in external)
    ledger_net = sum(l.amount for l in ledger if bank[0].time.date() <= l.time.date() <= bank[-1].time.date())
    bank_only_by_summary = {
        summary: [row for row in bank_only if row.summary == summary]
        for summary in sorted({row.summary for row in bank_only})
    }
    candidates = [
        (b, l) for b in bank_only for l in ledger_only
        if b.time.date() == l.time.date() and b.amount * l.amount > 0
        and abs((b.time - l.time).total_seconds()) <= 3600
        and abs(b.amount - l.amount) <= 100
    ]
    lines = [
        "# 兴业银行储蓄卡流水核对报告", "",
        f"- 银行范围：{bank[0].time:%Y-%m-%d} 至 {bank[-1].time:%Y-%m-%d}（UTC+8）。",
        f"- 原始 PDF：{len(bank)} 笔、{max(row.page for row in bank)} 页；线上账户：{account_name}；只读快照：{snapshot.name}。",
        f"- 银行收入：{yuan(income)} 元（{sum(b.amount > 0 for b in bank)} 笔）；支出：{yuan(-expense)} 元（{sum(b.amount < 0 for b in bank)} 笔）。",
        f"- 期初推导余额：{yuan(opening)} 元；期末余额：{yuan(bank[-1].balance)} 元；余额链错误：0 处。",
        f"- 卡内理财排除：{len(excluded)} 笔（理财代销、购汇），净额 {yuan(sum(b.amount for b in excluded))} 元；外部收支核对：{len(external)} 笔。",
        f"- 匹配：秒级同额 {counts['秒级同额']}、同日同额 {counts['同日同额']}、七日内同额 {counts['七日内同额']}；仅银行 {len(bank_only)}、仅线上 {len(ledger_only)}。",
        f"- 本期银行全量净变动：{yuan(sum(b.amount for b in bank))} 元；排除卡内理财后：{yuan(bank_net)} 元；线上同期净变动：{yuan(ledger_net)} 元；外部收支差额（银行－线上）：{yuan(bank_net - ledger_net)} 元。",
        "", "## 非秒级匹配", "",
        "| PDF位置 | 银行时间 | 金额 | 线上时间 | 时间差 | 交易ID |", "| --- | --- | ---: | --- | ---: | --- |",
    ]
    for b, l, phase in sorted(pairs, key=lambda pair: pair[0].time):
        if phase != "秒级同额":
            minutes = round((l.time - b.time).total_seconds() / 60, 1)
            lines.append(f"| p{b.page}/r{b.row} | {b.time:%Y-%m-%d %H:%M:%S} | {yuan(b.amount)} | {l.time:%Y-%m-%d %H:%M:%S} | {minutes:+g} 分钟 | {l.id} |")
    lines += ["", "## 排除的卡内理财交易", "", "| PDF位置 | 时间 | 金额 | 摘要 |", "| --- | --- | ---: | --- |"]
    for b in excluded:
        lines.append(f"| p{b.page}/r{b.row} | {b.time:%Y-%m-%d %H:%M:%S} | {yuan(b.amount)} | {b.summary} |")
    lines += ["", "## 差异分组", "", "| 银行摘要 | 笔数 | 净金额 |", "| --- | ---: | ---: |"]
    for summary, items in bank_only_by_summary.items():
        lines.append(f"| {summary} | {len(items)} | {yuan(sum(b.amount for b in items))} |")
    lines += ["", f"线上独有 {len(ledger_only)} 笔，净金额 {yuan(sum(l.amount for l in ledger_only))} 元。"]
    if account_balance is not None:
        lines.append(f"线上账户查询余额 {yuan(account_balance)} 元；与银行期末余额相差 {yuan(account_balance - bank[-1].balance)} 元。余额差额仍需结合账期和未匹配项目复核。")
    lines += ["", "## 同日金额差异候选（未自动匹配）", "",
              "| PDF位置 | 银行时间 | 银行金额 | 线上时间 | 线上金额 | 银行－线上 | 交易ID |",
              "| --- | --- | ---: | --- | ---: | ---: | --- |"]
    for b, l in candidates:
        lines.append(f"| p{b.page}/r{b.row} | {b.time:%Y-%m-%d %H:%M:%S} | {yuan(b.amount)} | {l.time:%Y-%m-%d %H:%M:%S} | {yuan(l.amount)} | {yuan(b.amount - l.amount)} | {l.id} |")
    lines += ["", "## 仅银行存在", "", "| PDF位置 | 时间 | 金额 | 摘要 |", "| --- | --- | ---: | --- |"]
    for b in bank_only:
        lines.append(f"| p{b.page}/r{b.row} | {b.time:%Y-%m-%d %H:%M:%S} | {yuan(b.amount)} | {b.summary} |")
    lines += ["", "## 仅线上存在", "", "| 时间 | 金额 | 类型 | 备注 | 交易ID |", "| --- | ---: | ---: | --- | --- |"]
    for l in ledger_only:
        lines.append(f"| {l.time:%Y-%m-%d %H:%M:%S} | {yuan(l.amount)} | {l.kind} | {l.comment} | {l.id} |")
    lines += ["", "## 口径与限制", "",
              "- PDF 表格按 9 pt 字形提取以排除跨格水印，Excel 往返读取后逐字段比对原 PDF。",
              "- 余额链包含全部银行流水；依用户确认，理财代销和购汇作为卡内理财从外部收支核对中排除，跨行代付仍参与匹配。",
              "- 线上交易按兴业储蓄卡视角取有符号账户金额；同额一对一匹配依次按秒、同日、前后 7 天进行。",
              "- 金额不同的记录保留在两侧差异中，不推断为同一交易。线上缓冲期内但银行范围外的记录不列入仅线上。",
              "- 线上快照通过认证令牌的 GET 接口获取；本核对脚本不调用业务写接口。", ""]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path)
    parser.add_argument("--online-json", type=Path)
    parser.add_argument("--account-id")
    parser.add_argument("--account-name", default="兴业银行储蓄卡")
    parser.add_argument("--account-balance-cents", type=int)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    if args.online_json and not args.account_id:
        raise ReconcileError("--account-id is required with --online-json")
    excel = args.pdf.with_suffix(".xlsx")
    if excel.exists() and not args.force:
        raise ReconcileError(f"Excel already exists: {excel}; pass --force to replace")
    bank = parse_pdf(args.pdf)
    write_excel(bank, args.pdf, excel)
    if read_excel(excel) != bank:
        raise ReconcileError("Excel round trip differs from PDF")
    print(f"Excel: {excel} | rows={len(bank)} | balance={yuan(bank[-1].balance)}")
    if args.online_json:
        ledger = read_ledger(args.online_json, args.account_id)
        output = args.output or args.pdf.with_name("兴业银行储蓄卡流水核对报告.md")
        output.write_text(report(bank, ledger, args.account_name, args.online_json,
                                 args.account_balance_cents), encoding="utf-8")
        print(f"Report: {output}")


if __name__ == "__main__":
    main()
