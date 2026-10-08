# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "pdfplumber==0.11.7",
#   "openpyxl==3.1.5",
# ]
# ///

"""Convert a CMB (招商银行) credit-card statement PDF into an auditable Excel workbook.

One PDF → one XLSX, one transaction row per PDF detail row. Conversion runs the
five statement-level balance checks and a round-trip validation against the PDF.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

from cmb_reconcile import (
    EXCEL_DETAIL_SHEET,
    EXCEL_HEADERS,
    LOCAL_ZONE,
    CATEGORY_CHARGE,
    CHARGE_CATEGORIES,
    CATEGORY_REFUND,
    CATEGORY_REPAYMENT,
    CmbStatement,
    CmbTransaction,
    ReconcileError,
    format_money,
    parse_statement,
    parse_statement_excel,
    validate_statement,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

SUMMARY_SHEET = "转换说明"
HEADER_FILL = PatternFill("solid", fgColor="C11B17")
SECTION_FILL = PatternFill("solid", fgColor="FCE4E4")
REFUND_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN_GRAY = Side(style="thin", color="E6D7D7")
MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00;0.00"
DATE_FORMAT = "yyyy-mm-dd"


def transaction_row(sequence: int, transaction: CmbTransaction) -> list[object]:
    return [
        sequence,
        transaction.trans_date,
        transaction.post_date,
        transaction.category,
        transaction.description,
        transaction.amount_minor / 100,
        transaction.card_last4,
        None if transaction.original_amount is None else float(transaction.original_amount),
        transaction.original_currency,
        transaction.page,
        transaction.row,
    ]


def write_statement_excel(
    statement: CmbStatement, output_path: Path
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = SUMMARY_SHEET
    details = workbook.create_sheet(EXCEL_DETAIL_SHEET)

    for column, header in enumerate(EXCEL_HEADERS, start=1):
        cell = details.cell(1, column, header)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN_GRAY)
    for sequence, transaction in enumerate(statement.transactions, start=1):
        details.append(transaction_row(sequence, transaction))

    last_row = len(statement.transactions) + 1
    table = Table(displayName="CMBStatement", ref=f"A1:K{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    details.add_table(table)
    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:K{last_row}"
    details.sheet_view.showGridLines = False
    details.row_dimensions[1].height = 26
    widths = {
        "A": 8,
        "B": 13,
        "C": 13,
        "D": 8,
        "E": 42,
        "F": 14,
        "G": 12,
        "H": 13,
        "I": 9,
        "J": 9,
        "K": 9,
    }
    for column, width in widths.items():
        details.column_dimensions[column].width = width
    for row in details.iter_rows(min_row=2, max_row=last_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
        row[1].number_format = DATE_FORMAT
        row[2].number_format = DATE_FORMAT
        row[5].number_format = MONEY_FORMAT
        row[7].number_format = MONEY_FORMAT
    details.conditional_formatting.add(
        f"F2:F{last_row}",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006")),
    )
    details.conditional_formatting.add(
        f"D2:D{last_row}",
        CellIsRule(operator="equal", formula=['"退款"'], fill=REFUND_FILL),
    )
    details.sheet_properties.pageSetUpPr.fitToPage = True
    details.page_setup.orientation = "landscape"
    details.page_setup.fitToWidth = 1
    details.page_setup.fitToHeight = 0
    details.print_title_rows = "1:1"

    summary.merge_cells("A1:C1")
    summary["A1"] = "招商银行信用卡对账单转换说明"
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = HEADER_FILL
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[1].height = 32
    source_hash = hashlib.sha256(statement.pdf_path.read_bytes()).hexdigest()
    metadata = [
        ("源文件", statement.pdf_path.name, "原始 PDF 文件名"),
        ("源文件 SHA-256", source_hash, "用于确认来源未变化"),
        ("转换时间", datetime.now(LOCAL_ZONE).strftime("%Y-%m-%d %H:%M:%S CST"), ""),
        ("时区", "UTC+8", "交易时间按中国标准时间解释，仅日期"),
    ]
    if statement.statement_date is not None:
        metadata.append(("账单日", statement.statement_date.strftime("%Y-%m-%d"), ""))
    if statement.new_balance_minor is not None:
        metadata.append(("本期应还金额", format_money(statement.new_balance_minor), ""))
    for row_index, values in enumerate(metadata, start=3):
        for column, value in enumerate(values, start=1):
            summary.cell(row_index, column, value)

    assumption_row = 3 + len(metadata) + 1
    summary.cell(assumption_row, 1, "处理约定")
    summary.cell(assumption_row, 1).font = Font(name="Arial", size=10, bold=True)
    summary.cell(assumption_row, 1).fill = SECTION_FILL
    assumptions = [
        "Excel 仅保存银行原始流水，一一对应 PDF 明细行，不合并、不拆分；后续核对只读取 Excel。",
        "消费记正、还款与退款记负；交易日与记账日均按 UTC+8 日期处理。",
        "还款行无记账日，记账日列留空，不虚构。",
        "原币金额与币种仅作追溯，不参与金额匹配。",
        "账单级余额等式（上期 + 流水合计 + 利息 = 本期应还）在转换时已校验通过。",
        "退款—消费抵销、未匹配补录等在后续核对阶段处理，本表不包含线上匹配结果。",
    ]
    for offset, value in enumerate(assumptions, start=1):
        summary.cell(assumption_row + offset, 1, f"{offset}. {value}")
        summary.merge_cells(
            start_row=assumption_row + offset,
            start_column=1,
            end_row=assumption_row + offset,
            end_column=3,
        )

    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 68
    summary.column_dimensions["C"].width = 34
    for row in summary.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Arial"
            cell.font = font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def convert_pdf_to_excel(
    pdf_path: Path, output_path: Path, *, overwrite: bool = False
) -> CmbStatement:
    pdf_path = pdf_path.resolve()
    output_path = output_path.resolve()
    if pdf_path.suffix.lower() != ".pdf" or not pdf_path.is_file():
        raise ReconcileError(f"input must be an existing PDF: {pdf_path}")
    if output_path.exists() and not overwrite:
        raise ReconcileError(
            f"Excel already exists; pass --force to overwrite: {output_path}"
        )
    password = os.environ.get("CMB_PDF_PASSWORD")
    statement = parse_statement(pdf_path, password)
    validate_statement(statement)
    write_statement_excel(statement, output_path)
    converted = parse_statement_excel(output_path)
    if [t for t in converted.transactions] != [t for t in statement.transactions]:
        raise ReconcileError("Excel round-trip validation did not match the PDF data")
    return statement


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path, help="CMB credit-card statement PDF")
    parser.add_argument("--output", type=Path, help="output XLSX path")
    parser.add_argument("--force", action="store_true", help="overwrite output XLSX")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    output = args.output or args.pdf.with_suffix(".xlsx")
    statement = convert_pdf_to_excel(args.pdf, output, overwrite=args.force)
    charges = sum(1 for t in statement.transactions if t.category in CHARGE_CATEGORIES)
    refunds = sum(1 for t in statement.transactions if t.category == CATEGORY_REFUND)
    repayments = sum(1 for t in statement.transactions if t.category == CATEGORY_REPAYMENT)
    print(
        f"Excel written to {output.resolve()} | rows={len(statement.transactions)} "
        f"charges={charges} refunds={refunds} repayments={repayments} "
        f"pages={max((t.page for t in statement.transactions), default=0)}"
    )
    if statement.new_balance_minor is not None:
        print(f"本期应还金额 ¥ {format_money(statement.new_balance_minor)}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ReconcileError as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1)
