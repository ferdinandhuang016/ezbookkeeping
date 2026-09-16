# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "pdfplumber==0.11.7",
#   "openpyxl==3.1.5",
# ]
# ///

"""Convert an ICBC historical-statement PDF into an auditable Excel workbook."""

from __future__ import annotations

import argparse
import hashlib
import os
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

from icbc_reconcile import (
    EXCEL_DETAIL_SHEET,
    EXCEL_HEADERS,
    LOCAL_ZONE,
    BankTransaction,
    ReconcileError,
    infer_pdf_password,
    parse_statement,
    parse_statement_excel,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

SUMMARY_SHEET = "转换说明"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
FINANCE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN_GRAY = Side(style="thin", color="D9E1F2")
MONEY_FORMAT = "#,##0.00;[Red]-#,##0.00;0.00"


def transaction_row(sequence: int, transaction: BankTransaction) -> list[object]:
    return [
        sequence,
        transaction.occurred_at.replace(tzinfo=None),
        transaction.account_number,
        transaction.deposit_type,
        transaction.serial_number,
        transaction.currency,
        transaction.cash_fx,
        transaction.summary,
        transaction.region,
        transaction.amount_minor / 100,
        transaction.balance_minor / 100,
        transaction.channel,
        transaction.page,
        transaction.row,
    ]


def write_statement_excel(
    transactions: list[BankTransaction], pdf_path: Path, output_path: Path
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = SUMMARY_SHEET
    details = workbook.create_sheet(EXCEL_DETAIL_SHEET)

    for column, header in enumerate(EXCEL_HEADERS, start=1):
        cell = details.cell(1, column, header)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN_GRAY)
    for sequence, transaction in enumerate(transactions, start=1):
        details.append(transaction_row(sequence, transaction))

    last_row = len(transactions) + 1
    table = Table(displayName="ICBCStatement", ref=f"A1:N{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    details.add_table(table)
    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:N{last_row}"
    details.sheet_view.showGridLines = False
    details.row_dimensions[1].height = 26
    widths = {
        "A": 8,
        "B": 21,
        "C": 23,
        "D": 13,
        "E": 25,
        "F": 10,
        "G": 11,
        "H": 18,
        "I": 11,
        "J": 17,
        "K": 17,
        "L": 15,
        "M": 10,
        "N": 10,
    }
    for column, width in widths.items():
        details.column_dimensions[column].width = width
    for row in details.iter_rows(min_row=2, max_row=last_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
        row[1].number_format = "yyyy-mm-dd hh:mm:ss"
        row[9].number_format = MONEY_FORMAT
        row[10].number_format = MONEY_FORMAT
    details.conditional_formatting.add(
        f"J2:J{last_row}",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006")),
    )
    details.conditional_formatting.add(
        f"H2:H{last_row}",
        CellIsRule(
            operator="equal",
            formula=['"理财"'],
            fill=FINANCE_FILL,
        ),
    )
    details.sheet_properties.pageSetUpPr.fitToPage = True
    details.page_setup.orientation = "landscape"
    details.page_setup.fitToWidth = 1
    details.page_setup.fitToHeight = 0
    details.print_title_rows = "1:1"

    summary.merge_cells("A1:C1")
    summary["A1"] = "工商银行历史明细转换说明"
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = HEADER_FILL
    summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[1].height = 32
    source_hash = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    metadata = [
        ("源文件", pdf_path.name, "原始 PDF 文件名"),
        ("源文件 SHA-256", source_hash, "用于确认来源未变化"),
        ("转换时间", datetime.now(LOCAL_ZONE).strftime("%Y-%m-%d %H:%M:%S CST"), ""),
        ("时区", "UTC+8", "交易时间按中国标准时间解释"),
    ]
    for row_index, values in enumerate(metadata, start=3):
        for column, value in enumerate(values, start=1):
            summary.cell(row_index, column, value)

    assumption_row = 8
    summary.cell(assumption_row, 1, "处理约定")
    summary.cell(assumption_row, 1).font = Font(name="Arial", size=10, bold=True)
    summary.cell(assumption_row, 1).fill = SECTION_FILL
    assumptions = [
        "Excel 仅保存银行原始流水，不包含线上匹配结果或待执行修改。",
        "理财和金融付款的归类、同日匹配及 2.60–2.80 CNY 地铁规则在后续核对阶段处理。",
        "PDF 页码和行号用于追溯；不得只凭手工排序后的 Excel 行号定位原始流水。",
    ]
    for offset, value in enumerate(assumptions, start=1):
        summary.cell(assumption_row + offset, 1, f"{offset}. {value}")
        summary.merge_cells(
            start_row=assumption_row + offset,
            start_column=1,
            end_row=assumption_row + offset,
            end_column=3,
        )

    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 68
    summary.column_dimensions["C"].width = 34
    for row in summary.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Arial"
            cell.font = font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def convert_pdf_to_excel(
    pdf_path: Path, output_path: Path, *, overwrite: bool = False
) -> list[BankTransaction]:
    pdf_path = pdf_path.resolve()
    output_path = output_path.resolve()
    if pdf_path.suffix.lower() != ".pdf" or not pdf_path.is_file():
        raise ReconcileError(f"input must be an existing PDF: {pdf_path}")
    if output_path.exists() and not overwrite:
        raise ReconcileError(
            f"Excel already exists; pass --force to overwrite: {output_path}"
        )
    password = os.environ.get("ICBC_PDF_PASSWORD") or infer_pdf_password(pdf_path)
    transactions = parse_statement(pdf_path, password)
    write_statement_excel(transactions, pdf_path, output_path)
    converted = parse_statement_excel(output_path)
    if converted != transactions:
        raise ReconcileError("Excel round-trip validation did not match the PDF data")
    return transactions


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path, help="ICBC historical-statement PDF")
    parser.add_argument("--output", type=Path, help="output XLSX path")
    parser.add_argument("--force", action="store_true", help="overwrite output XLSX")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    output = args.output or args.pdf.with_suffix(".xlsx")
    transactions = convert_pdf_to_excel(args.pdf, output, overwrite=args.force)
    print(
        f"Excel written to {output.resolve()} | rows={len(transactions)} "
        f"pages={max(item.page for item in transactions)}"
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ReconcileError as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1)
