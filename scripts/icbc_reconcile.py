# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "pdfplumber==0.11.7",
#   "openpyxl==3.1.5",
# ]
# ///

"""Reconcile an ICBC debit-card statement with ezBookkeeping, read-only."""

from __future__ import annotations

import argparse
import getpass
import json
import os
import re
import ssl
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from itertools import pairwise
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook

LOCAL_ZONE = timezone(timedelta(hours=8), "CST")
DATETIME_FORMAT = "%Y-%m-%d\n%H:%M:%S"
MONEY_PATTERN = re.compile(r"^[+-]?(?:\d{1,3}(?:,\d{3})*|\d+)\.\d{2}$")
FINANCE_SUMMARIES = {"理财", "金融付款"}
DAILY_FINANCE_PREFIX = "工商银行理财按日拆分"
PRE_STATEMENT_RESIDUAL_COMMENT = "工商银行账单期前理财残差"


class ReconcileError(RuntimeError):
    """Raised when input or remote data cannot be reconciled safely."""


@dataclass(frozen=True)
class BankTransaction:
    occurred_at: datetime
    amount_minor: int
    balance_minor: int
    page: int
    row: int
    summary: str = ""
    account_number: str = ""
    deposit_type: str = ""
    serial_number: str = ""
    currency: str = ""
    cash_fx: str = ""
    region: str = ""
    channel: str = ""

    @property
    def key(self) -> tuple[int, int]:
        return int(self.occurred_at.timestamp()), self.amount_minor


@dataclass(frozen=True)
class OnlineTransaction:
    occurred_at: datetime
    amount_minor: int
    transaction_id: str
    transaction_type: int
    comment: str
    source_account_id: str = ""
    destination_account_id: str = ""

    @property
    def key(self) -> tuple[int, int]:
        return int(self.occurred_at.timestamp()), self.amount_minor


@dataclass(frozen=True)
class NearMatch:
    bank: BankTransaction
    online: OnlineTransaction

    @property
    def time_delta_seconds(self) -> int:
        return abs(self.bank.key[0] - self.online.key[0])

    @property
    def amount_delta_minor(self) -> int:
        return self.bank.amount_minor - self.online.amount_minor


@dataclass(frozen=True)
class Reconciliation:
    exact_matched_count: int
    near_matches: list[NearMatch]
    bank_only: list[BankTransaction]
    online_only: list[OnlineTransaction]

    @property
    def matched_count(self) -> int:
        return self.exact_matched_count + len(self.near_matches)


@dataclass(frozen=True)
class BankOffsetPair:
    outflow: BankTransaction
    inflow: BankTransaction


@dataclass(frozen=True)
class FinanceAggregation:
    online: OnlineTransaction
    bank_transactions: list[BankTransaction]
    period_start: datetime
    partial_start: bool = False

    @property
    def bank_expense_minor(self) -> int:
        return -sum(transaction.amount_minor for transaction in self.bank_transactions)

    @property
    def online_expense_minor(self) -> int:
        return -self.online.amount_minor

    @property
    def expense_delta_minor(self) -> int:
        return self.bank_expense_minor - self.online_expense_minor


@dataclass(frozen=True)
class FinanceReconciliation:
    aggregations: list[FinanceAggregation]
    pending_bank: list[BankTransaction]
    regular_bank: list[BankTransaction]
    regular_online: list[OnlineTransaction]
    excluded_residual_online: list[OnlineTransaction]


def parse_money(value: str, *, location: str) -> int:
    normalized = value.strip().replace(" ", "")
    if not MONEY_PATTERN.fullmatch(normalized):
        raise ReconcileError(f"{location}: invalid money value {value!r}")
    try:
        return int(Decimal(normalized.replace(",", "")) * 100)
    except (InvalidOperation, ValueError) as exc:
        raise ReconcileError(f"{location}: invalid money value {value!r}") from exc


def infer_pdf_password(pdf_path: Path) -> str | None:
    match = re.search(r"_(\d{6})$", pdf_path.stem)
    return match.group(1) if match else None


def clean_cell(value: Any) -> str:
    return " ".join(str(value or "").split())


def parse_statement(pdf_path: Path, password: str | None) -> list[BankTransaction]:
    try:
        pdf = pdfplumber.open(pdf_path, password=password)
    except Exception as exc:
        raise ReconcileError(f"cannot open PDF {pdf_path}: {exc}") from exc

    transactions: list[BankTransaction] = []
    with pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            body = page.filter(
                lambda obj: (
                    obj.get("object_type") != "char"
                    or abs(float(obj.get("size", 0)) - 7.0) < 0.05
                )
            )
            tables = body.extract_tables()
            if len(tables) != 1 or not tables[0]:
                raise ReconcileError(
                    f"page {page_number}: expected one transaction table, found {len(tables)}"
                )

            rows = tables[0]
            for row_number, row in enumerate(rows[1:], start=1):
                if len(row) != 11:
                    raise ReconcileError(
                        f"page {page_number}, row {row_number}: expected 11 columns, found {len(row)}"
                    )
                location = f"page {page_number}, row {row_number}"
                try:
                    occurred_at = datetime.strptime(
                        (row[0] or "").strip(), DATETIME_FORMAT
                    ).replace(tzinfo=LOCAL_ZONE)
                except ValueError as exc:
                    raise ReconcileError(
                        f"{location}: invalid transaction datetime {row[0]!r}"
                    ) from exc

                transactions.append(
                    BankTransaction(
                        occurred_at=occurred_at,
                        amount_minor=parse_money(
                            row[8] or "", location=f"{location} amount"
                        ),
                        balance_minor=parse_money(
                            row[9] or "", location=f"{location} balance"
                        ),
                        page=page_number,
                        row=row_number,
                        summary=clean_cell(row[6]),
                        account_number=clean_cell(row[1]),
                        deposit_type=clean_cell(row[2]),
                        serial_number=clean_cell(row[3]),
                        currency=clean_cell(row[4]),
                        cash_fx=clean_cell(row[5]),
                        region=clean_cell(row[7]),
                        channel=clean_cell(row[10]),
                    )
                )

    validate_statement(transactions)
    return transactions


EXCEL_DETAIL_SHEET = "工商银行流水"
EXCEL_HEADERS = (
    "序号",
    "交易时间",
    "账号",
    "存款种类",
    "交易流水号",
    "币种",
    "钞汇标志",
    "摘要",
    "地区",
    "收支金额（元）",
    "账户余额（元）",
    "交易渠道",
    "PDF页码",
    "PDF行号",
)


def excel_money_to_minor(value: Any, *, location: str) -> int:
    try:
        decimal_value = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ReconcileError(
            f"{location}: invalid Excel money value {value!r}"
        ) from exc
    scaled = decimal_value * 100
    if scaled != scaled.to_integral_value():
        raise ReconcileError(f"{location}: money has more than two decimals {value!r}")
    return int(scaled)


def parse_statement_excel(excel_path: Path) -> list[BankTransaction]:
    try:
        workbook = load_workbook(excel_path, read_only=True, data_only=False)
    except Exception as exc:
        raise ReconcileError(f"cannot open Excel {excel_path}: {exc}") from exc

    try:
        if EXCEL_DETAIL_SHEET not in workbook.sheetnames:
            raise ReconcileError(
                f"Excel is missing required sheet {EXCEL_DETAIL_SHEET!r}"
            )
        sheet = workbook[EXCEL_DETAIL_SHEET]
        headers = tuple(
            cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        )
        if headers != EXCEL_HEADERS:
            raise ReconcileError("Excel detail headers do not match the ICBC schema")

        transactions: list[BankTransaction] = []
        for excel_row, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if not any(value is not None for value in values):
                continue
            if len(values) != len(EXCEL_HEADERS):
                raise ReconcileError(f"Excel row {excel_row}: invalid column count")
            expected_sequence = len(transactions) + 1
            if values[0] != expected_sequence:
                raise ReconcileError(
                    f"Excel row {excel_row}: expected sequence {expected_sequence}"
                )
            occurred_at = values[1]
            if isinstance(occurred_at, str):
                try:
                    occurred_at = datetime.strptime(
                        occurred_at, "%Y-%m-%d %H:%M:%S"
                    ).replace(tzinfo=LOCAL_ZONE)
                except ValueError as exc:
                    raise ReconcileError(
                        f"Excel row {excel_row}: invalid transaction datetime"
                    ) from exc
            if not isinstance(occurred_at, datetime):
                raise ReconcileError(
                    f"Excel row {excel_row}: transaction time is not a datetime"
                )
            if occurred_at.tzinfo is None:
                occurred_at = occurred_at.replace(tzinfo=LOCAL_ZONE)
            else:
                occurred_at = occurred_at.astimezone(LOCAL_ZONE)
            transactions.append(
                BankTransaction(
                    occurred_at=occurred_at,
                    amount_minor=excel_money_to_minor(
                        values[9], location=f"Excel row {excel_row} amount"
                    ),
                    balance_minor=excel_money_to_minor(
                        values[10], location=f"Excel row {excel_row} balance"
                    ),
                    page=int(values[12]),
                    row=int(values[13]),
                    summary=clean_cell(values[7]),
                    account_number=clean_cell(values[2]),
                    deposit_type=clean_cell(values[3]),
                    serial_number=clean_cell(values[4]),
                    currency=clean_cell(values[5]),
                    cash_fx=clean_cell(values[6]),
                    region=clean_cell(values[8]),
                    channel=clean_cell(values[11]),
                )
            )
    finally:
        workbook.close()

    validate_statement(transactions)
    return transactions


def load_statement(statement_path: Path) -> list[BankTransaction]:
    suffix = statement_path.suffix.lower()
    if suffix == ".pdf":
        password = os.environ.get("ICBC_PDF_PASSWORD") or infer_pdf_password(
            statement_path
        )
        return parse_statement(statement_path, password)
    if suffix == ".xlsx":
        return parse_statement_excel(statement_path)
    raise ReconcileError("statement must be an ICBC .pdf or converted .xlsx file")


def validate_statement(transactions: list[BankTransaction]) -> None:
    if not transactions:
        raise ReconcileError("the statement contains no transactions")

    for previous, current in pairwise(transactions):
        if current.occurred_at < previous.occurred_at:
            raise ReconcileError(
                f"page {current.page}, row {current.row}: transactions are not chronological"
            )
        expected_balance = previous.balance_minor + current.amount_minor
        if current.balance_minor != expected_balance:
            raise ReconcileError(
                f"page {current.page}, row {current.row}: running balance mismatch; "
                f"expected {format_money(expected_balance)}, got {format_money(current.balance_minor)}"
            )


class EzBookkeepingClient:
    """Minimal client that blocks business-data writes unless explicitly enabled."""

    def __init__(
        self,
        base_url: str,
        timeout: float = 30.0,
        allowed_business_posts: frozenset[str] = frozenset(),
    ) -> None:
        parsed = urllib.parse.urlsplit(base_url)
        if parsed.scheme != "https" or not parsed.netloc:
            raise ReconcileError("base URL must be an absolute HTTPS URL")
        path = parsed.path.rstrip("/")
        self.base_url = urllib.parse.urlunsplit(
            (parsed.scheme, parsed.netloc, f"{path}/api", "", "")
            if not path.endswith("/api")
            else (parsed.scheme, parsed.netloc, path, "", "")
        )
        self.timeout = timeout
        self.token: str | None = None
        self.ssl_context = ssl.create_default_context()
        self.allowed_business_posts = allowed_business_posts

    def authorize(self, login_name: str, password: str) -> None:
        result = self._request(
            "POST",
            "/authorize.json",
            {"loginName": login_name, "password": password},
        )
        if result.get("need2FA"):
            raise ReconcileError(
                "this account requires 2FA; password-only reconciliation cannot continue"
            )
        token = result.get("token")
        if not isinstance(token, str) or not token:
            raise ReconcileError("authorization response did not contain a token")
        self.token = token

    def list_accounts(self) -> list[dict[str, Any]]:
        result = self._get("/v1/accounts/list.json", {"visible_only": "false"})
        if not isinstance(result, list):
            raise ReconcileError("account list response is not a list")
        return list(flatten_accounts(result))

    def list_transactions(
        self,
        account_id: str,
        start_time: int,
        end_time: int,
        *,
        with_pictures: bool = False,
    ) -> list[dict[str, Any]]:
        result = self._get(
            "/v1/transactions/list/all.json",
            {
                "account_ids": account_id,
                "start_time": start_time,
                "end_time": end_time,
                "with_pictures": "true" if with_pictures else "false",
                "trim_account": "true",
                "trim_category": "true",
                "trim_tag": "true",
            },
        )
        if not isinstance(result, list):
            raise ReconcileError("transaction list response is not a list")
        return result

    def _get(self, path: str, query: dict[str, str | int]) -> Any:
        encoded = urllib.parse.urlencode(query)
        return self._request("GET", f"{path}?{encoded}")

    def _request(
        self, method: str, path: str, data: dict[str, Any] | None = None
    ) -> Any:
        if method not in {"GET", "POST"} or (
            method == "POST"
            and path != "/authorize.json"
            and path not in self.allowed_business_posts
        ):
            raise ReconcileError(f"blocked non-read-only request: {method} {path}")

        body = json.dumps(data).encode("utf-8") if data is not None else None
        headers = {
            "Accept": "application/json",
            "Accept-Language": "zh-Hans",
            "X-Timezone-Name": "Asia/Shanghai",
            "X-Timezone-Offset": "-480",
        }
        if body is not None:
            headers["Content-Type"] = "application/json"
        if self.token:
            headers["Authorization"] = f"Bearer {self.token}"

        request = urllib.request.Request(
            f"{self.base_url}{path}", data=body, headers=headers, method=method
        )
        try:
            with urllib.request.urlopen(
                request, timeout=self.timeout, context=self.ssl_context
            ) as response:
                payload = json.load(response)
        except urllib.error.HTTPError as exc:
            message = exc.read().decode("utf-8", errors="replace")
            raise ReconcileError(f"HTTP {exc.code} for {path}: {message}") from exc
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as exc:
            raise ReconcileError(f"request failed for {path}: {exc}") from exc

        if not isinstance(payload, dict) or not payload.get("success"):
            code = (
                payload.get("errorCode", "unknown")
                if isinstance(payload, dict)
                else "unknown"
            )
            message = (
                payload.get("errorMessage", "invalid response")
                if isinstance(payload, dict)
                else "invalid response"
            )
            raise ReconcileError(f"API error {code} for {path}: {message}")
        return payload.get("result")


def flatten_accounts(accounts: Iterable[dict[str, Any]]) -> Iterable[dict[str, Any]]:
    for account in accounts:
        yield account
        children = account.get("subAccounts") or []
        if isinstance(children, list):
            yield from flatten_accounts(children)


def select_account(
    accounts: list[dict[str, Any]], selector: str | None, closing_balance: int
) -> dict[str, Any]:
    if selector:
        matches = [
            account
            for account in accounts
            if str(account.get("id")) == selector
            or str(account.get("name")) == selector
        ]
    else:
        matches = [
            account
            for account in accounts
            if account.get("currency") == "CNY"
            and int(account.get("balance", 0)) == closing_balance
        ]

    if len(matches) == 1:
        return matches[0]

    choices = ", ".join(
        f"{account.get('name')} ({account.get('id')}, {account.get('currency')}, "
        f"{format_money(int(account.get('balance', 0)))})"
        for account in accounts
    )
    reason = "did not match" if not matches else "was ambiguous"
    raise ReconcileError(
        f"account selection {reason}; pass --account with an exact name or ID. Available: {choices}"
    )


def normalize_online_transactions(
    rows: list[dict[str, Any]], account_id: str
) -> tuple[list[OnlineTransaction], list[dict[str, Any]]]:
    normalized: list[OnlineTransaction] = []
    unsupported: list[dict[str, Any]] = []
    for row in rows:
        transaction_type = int(row.get("type", 0))
        source_account_id = str(row.get("sourceAccountId") or "")
        destination_account_id = str(row.get("destinationAccountId") or "")

        if transaction_type == 2:
            amount_minor = int(row["sourceAmount"])
        elif (
            transaction_type == 3
            or transaction_type == 4
            and source_account_id == account_id
        ):
            amount_minor = -int(row["sourceAmount"])
        elif transaction_type == 4 and destination_account_id == account_id:
            amount_minor = int(row["destinationAmount"])
        elif transaction_type == 1 and row.get("balanceDelta") is not None:
            amount_minor = int(row["balanceDelta"])
        else:
            unsupported.append(row)
            continue

        normalized.append(
            OnlineTransaction(
                occurred_at=datetime.fromtimestamp(int(row["time"]), LOCAL_ZONE),
                amount_minor=amount_minor,
                transaction_id=str(row.get("id") or ""),
                transaction_type=transaction_type,
                comment=str(row.get("comment") or ""),
                source_account_id=source_account_id,
                destination_account_id=destination_account_id,
            )
        )
    return normalized, unsupported


def partition_online_by_period(
    transactions: list[OnlineTransaction],
    period_start: datetime,
    period_end: datetime,
    tolerance_seconds: int,
) -> tuple[list[OnlineTransaction], list[OnlineTransaction]]:
    if tolerance_seconds < 0:
        raise ReconcileError("time tolerance must not be negative")
    boundary_tolerance = timedelta(seconds=tolerance_seconds)
    lower_bound = period_start - boundary_tolerance
    upper_bound = period_end + boundary_tolerance
    within: list[OnlineTransaction] = []
    outside: list[OnlineTransaction] = []
    for transaction in transactions:
        target = (
            within if lower_bound <= transaction.occurred_at <= upper_bound else outside
        )
        target.append(transaction)
    return within, outside


def reconcile_finance_aggregations(
    bank_transactions: list[BankTransaction],
    online_transactions: list[OnlineTransaction],
    account_id: str,
    fund_account_ids: set[str],
) -> FinanceReconciliation:
    finance_bank = sorted(
        (
            transaction
            for transaction in bank_transactions
            if transaction.summary.strip() in FINANCE_SUMMARIES
            and transaction.amount_minor < 0
        ),
        key=lambda transaction: transaction.occurred_at,
    )
    residual_online = [
        transaction
        for transaction in online_transactions
        if transaction.transaction_type == 4
        and transaction.source_account_id == account_id
        and transaction.destination_account_id in fund_account_ids
        and transaction.comment.strip() == PRE_STATEMENT_RESIDUAL_COMMENT
    ]
    finance_online = sorted(
        (
            transaction
            for transaction in online_transactions
            if transaction.transaction_type == 4
            and transaction.source_account_id == account_id
            and transaction.destination_account_id in fund_account_ids
            and transaction.comment.strip() != PRE_STATEMENT_RESIDUAL_COMMENT
        ),
        key=lambda transaction: transaction.occurred_at,
    )

    statement_start = bank_transactions[0].occurred_at
    previous_transfer_time = statement_start
    bank_index = 0
    aggregations: list[FinanceAggregation] = []
    for transfer_index, transfer in enumerate(finance_online):
        period_transactions: list[BankTransaction] = []
        while (
            bank_index < len(finance_bank)
            and finance_bank[bank_index].occurred_at <= transfer.occurred_at
        ):
            transaction = finance_bank[bank_index]
            if transaction.occurred_at > previous_transfer_time or transfer_index == 0:
                period_transactions.append(transaction)
            bank_index += 1
        aggregations.append(
            FinanceAggregation(
                online=transfer,
                bank_transactions=period_transactions,
                period_start=previous_transfer_time,
                partial_start=transfer_index == 0
                and not transfer.comment.startswith(DAILY_FINANCE_PREFIX),
            )
        )
        previous_transfer_time = transfer.occurred_at

    pending_bank = finance_bank[bank_index:]
    finance_bank_ids = {id(transaction) for transaction in finance_bank}
    finance_online_ids = {
        id(transaction) for transaction in finance_online + residual_online
    }
    return FinanceReconciliation(
        aggregations=aggregations,
        pending_bank=pending_bank,
        regular_bank=[
            transaction
            for transaction in bank_transactions
            if id(transaction) not in finance_bank_ids
        ],
        regular_online=[
            transaction
            for transaction in online_transactions
            if id(transaction) not in finance_online_ids
        ],
        excluded_residual_online=residual_online,
    )


def reconcile(
    bank_transactions: list[BankTransaction],
    online_transactions: list[OnlineTransaction],
    amount_tolerance_minor: int = 5,
    small_amount_limit_minor: int = 1000,
) -> Reconciliation:
    if amount_tolerance_minor < 0 or small_amount_limit_minor < 0:
        raise ReconcileError("matching tolerances cannot be negative")

    online_by_key: dict[tuple[int, int], list[int]] = defaultdict(list)
    for index, transaction in enumerate(online_transactions):
        online_by_key[transaction.key].append(index)

    matched_bank: set[int] = set()
    matched_online: set[int] = set()
    for bank_index, transaction in enumerate(bank_transactions):
        candidates = online_by_key.get(transaction.key)
        if not candidates:
            continue
        while candidates and candidates[-1] in matched_online:
            candidates.pop()
        if candidates:
            matched_bank.add(bank_index)
            matched_online.add(candidates.pop())

    fuzzy_candidates: list[tuple[int, int, int, int]] = []
    for bank_index, bank in enumerate(bank_transactions):
        if bank_index in matched_bank:
            continue
        for online_index, online in enumerate(online_transactions):
            if online_index in matched_online:
                continue
            time_delta = abs(bank.key[0] - online.key[0])
            amount_delta = abs(bank.amount_minor - online.amount_minor)
            normal_candidate = amount_delta <= amount_tolerance_minor
            small_candidate = (
                abs(bank.amount_minor) <= small_amount_limit_minor
                and abs(online.amount_minor) <= small_amount_limit_minor
            )
            same_direction = bank.amount_minor * online.amount_minor > 0
            if (
                same_direction
                and bank.occurred_at.date() == online.occurred_at.date()
                and (normal_candidate or small_candidate)
            ):
                fuzzy_candidates.append(
                    (time_delta, amount_delta, bank_index, online_index)
                )

    near_matches: list[NearMatch] = []
    for _, _, bank_index, online_index in sorted(fuzzy_candidates):
        if bank_index in matched_bank or online_index in matched_online:
            continue
        matched_bank.add(bank_index)
        matched_online.add(online_index)
        near_matches.append(
            NearMatch(bank_transactions[bank_index], online_transactions[online_index])
        )

    return Reconciliation(
        exact_matched_count=len(matched_bank) - len(near_matches),
        near_matches=sorted(near_matches, key=lambda item: item.bank.occurred_at),
        bank_only=[
            transaction
            for index, transaction in enumerate(bank_transactions)
            if index not in matched_bank
        ],
        online_only=[
            transaction
            for index, transaction in enumerate(online_transactions)
            if index not in matched_online
        ],
    )


def suppress_equal_opposite_bank_only(
    transactions: list[BankTransaction],
) -> tuple[list[BankTransaction], list[BankOffsetPair]]:
    candidates: list[tuple[int, int, int]] = []
    for left_index, left in enumerate(transactions):
        for right_index in range(left_index + 1, len(transactions)):
            right = transactions[right_index]
            if left.amount_minor + right.amount_minor == 0:
                candidates.append(
                    (abs(left.key[0] - right.key[0]), left_index, right_index)
                )

    suppressed: set[int] = set()
    pairs: list[BankOffsetPair] = []
    for _, left_index, right_index in sorted(candidates):
        if left_index in suppressed or right_index in suppressed:
            continue
        left = transactions[left_index]
        right = transactions[right_index]
        outflow, inflow = (left, right) if left.amount_minor < 0 else (right, left)
        pairs.append(BankOffsetPair(outflow, inflow))
        suppressed.update((left_index, right_index))

    remaining = [
        transaction
        for index, transaction in enumerate(transactions)
        if index not in suppressed
    ]
    return remaining, sorted(pairs, key=lambda item: item.outflow.occurred_at)


def format_money(amount_minor: int) -> str:
    sign = "+" if amount_minor > 0 else "-" if amount_minor < 0 else ""
    absolute = abs(amount_minor)
    return f"{sign}{absolute // 100:,}.{absolute % 100:02d}"


def summarize(transactions: Iterable[Any]) -> tuple[int, int, int]:
    values = [transaction.amount_minor for transaction in transactions]
    income = sum(value for value in values if value > 0)
    expense = -sum(value for value in values if value < 0)
    return len(values), income, expense


def daily_summary(transactions: Iterable[Any]) -> dict[str, tuple[int, int, int]]:
    grouped: dict[str, list[Any]] = defaultdict(list)
    for transaction in transactions:
        grouped[transaction.occurred_at.date().isoformat()].append(transaction)
    return {day: summarize(items) for day, items in grouped.items()}


def render_report(
    bank_transactions: list[BankTransaction],
    online_transactions: list[OnlineTransaction],
    unsupported: list[dict[str, Any]],
    finance: FinanceReconciliation,
    result: Reconciliation,
    suppressed_bank_pairs: list[BankOffsetPair],
    account: dict[str, Any],
    base_url: str,
    detail_limit: int,
    time_tolerance_seconds: int,
    amount_tolerance_minor: int,
    small_amount_limit_minor: int,
    account_selector: str | None,
    fund_account_name: str,
    outside_period_online_count: int,
) -> str:
    bank_count, bank_income, bank_expense = summarize(bank_transactions)
    online_count, online_income, online_expense = summarize(online_transactions)
    opening_balance = (
        bank_transactions[0].balance_minor - bank_transactions[0].amount_minor
    )
    closing_balance = bank_transactions[-1].balance_minor
    online_balance = int(account.get("balance", 0))
    bank_daily = daily_summary(bank_transactions)
    online_daily = daily_summary(online_transactions)
    near_same_amount = sum(
        match.amount_delta_minor == 0 for match in result.near_matches
    )
    near_amount_difference = len(result.near_matches) - near_same_amount
    near_amount_delta = sum(match.amount_delta_minor for match in result.near_matches)
    amount_difference_matches = [
        match for match in result.near_matches if match.amount_delta_minor != 0
    ]
    allocated_finance_bank = [
        transaction
        for aggregation in finance.aggregations
        for transaction in aggregation.bank_transactions
    ]
    allocated_finance_expense = -sum(
        transaction.amount_minor for transaction in allocated_finance_bank
    )
    online_finance_expense = sum(
        aggregation.online_expense_minor for aggregation in finance.aggregations
    )
    complete_finance = [
        aggregation
        for aggregation in finance.aggregations
        if not aggregation.partial_start
    ]
    complete_finance_bank_expense = sum(
        aggregation.bank_expense_minor for aggregation in complete_finance
    )
    complete_online_finance_expense = sum(
        aggregation.online_expense_minor for aggregation in complete_finance
    )
    partial_finance = [
        aggregation for aggregation in finance.aggregations if aggregation.partial_start
    ]
    partial_finance_delta = sum(
        aggregation.expense_delta_minor for aggregation in partial_finance
    )
    pending_finance_expense = -sum(
        transaction.amount_minor for transaction in finance.pending_bank
    )
    residual_finance_expense = -sum(
        transaction.amount_minor for transaction in finance.excluded_residual_online
    )
    full_scope_finance_expense = allocated_finance_expense + pending_finance_expense
    partial_summary = (
        f"首个不完整区间的金额差为 {format_money(partial_finance_delta)} CNY，仅表示 PDF 覆盖不足；"
        if partial_finance
        else "理财流水已按日拆分，所有区间都位于 PDF 覆盖范围内；"
    )

    lines = [
        "# 工商银行收支核对报告",
        "",
        f"- 生成时间：{datetime.now(LOCAL_ZONE).strftime('%Y-%m-%d %H:%M:%S %Z')}",
        f"- 环境：{base_url.rstrip('/')}",
        f"- 线上账户：{account.get('name')}（ID `{account.get('id')}`，CNY）",
        f"- 账户选择：{'显式指定' if account_selector else '按 PDF 期末余额唯一匹配'}",
        f"- 银行流水范围：{bank_transactions[0].occurred_at:%Y-%m-%d %H:%M:%S} 至 {bank_transactions[-1].occurred_at:%Y-%m-%d %H:%M:%S}",
        f"- 线上范围过滤：排除 PDF 精确起止时间（边界容差 {time_tolerance_seconds} 秒）之外的 {outside_period_online_count} 笔非理财流水。",
        "- 操作范围：仅登录并调用账户列表、交易列表 GET 接口；未新增、修改或删除任何记账记录。认证可能刷新用户最后登录时间。",
        "",
        "## 结论",
        "",
        (
            f"排除汇总理财流水后，精确匹配 {result.exact_matched_count} 笔，"
            f"容差匹配 {len(result.near_matches)} 笔；银行有但线上没有 {len(result.bank_only)} 笔；"
            f"线上有但银行没有 {len(result.online_only)} 笔。"
        ),
        f"另有 {len(suppressed_bank_pairs)} 组等额转入/转出银行流水互相抵销，按规则不计入未匹配统计。",
        (
            f"容差匹配中 {near_same_amount} 笔仅时间有偏差，{near_amount_difference} 笔同时有金额偏差；"
            f"这些金额差（银行 - 线上）合计 {format_money(near_amount_delta)} CNY。"
        ),
        (
            f"另有 {len(finance.aggregations)} 笔转入“{fund_account_name}”的线上转账，"
            f"对应 {len(allocated_finance_bank)} 笔 PDF“理财/金融付款”流水；"
            f"其中 {len(complete_finance)} 个完整区间金额差为 "
            f"{format_money(complete_finance_bank_expense - complete_online_finance_expense)} CNY。"
        ),
        partial_summary
        + f"末次汇总后还有 {len(finance.pending_bank)} 笔、{format_money(pending_finance_expense)} CNY 待汇总。",
        f"另排除 {len(finance.excluded_residual_online)} 笔、{format_money(residual_finance_expense)} CNY 的账单期前理财残差。",
        f"PDF 期末余额为 {format_money(closing_balance)} CNY，线上账户当前余额为 {format_money(online_balance)} CNY，"
        + ("两者一致。" if closing_balance == online_balance else "两者不一致。"),
        "",
        "## 汇总",
        "",
        "| 来源 | 笔数 | 收入 | 支出 | 净额 |",
        "|---|---:|---:|---:|---:|",
        f"| 银行 PDF | {bank_count} | {format_money(bank_income)} | {format_money(bank_expense)} | {format_money(bank_income - bank_expense)} |",
        f"| 线上账本 | {online_count} | {format_money(online_income)} | {format_money(online_expense)} | {format_money(online_income - online_expense)} |",
        f"| 差异（银行 - 线上） | {bank_count - online_count:+d} | {format_money(bank_income - online_income)} | {format_money(bank_expense - online_expense)} | {format_money((bank_income - bank_expense) - (online_income - online_expense))} |",
        "",
        f"- PDF 期初余额：{format_money(opening_balance)} CNY",
        f"- PDF 期末余额：{format_money(closing_balance)} CNY",
        f"- 无法安全换算的线上余额调整记录：{len(unsupported)} 笔",
        "",
        "## 理财汇总区间差异",
        "",
        "PDF 摘要为“理财”或“金融付款”的支出，按相邻两次线上转入支付宝基金的时间切分；区间为前开后闭。首个区间可能包含 PDF 起始日前的流水，因此标记为不完整。",
        "",
        "| 汇总截止时间 | PDF 理财（笔/金额） | PDF 金融付款（笔/金额） | PDF 合计 | 线上汇总转账 | 差额（PDF - 线上） | 备注 |",
        "|---|---:|---:|---:|---:|---:|---|",
    ]

    for aggregation in finance.aggregations:
        wealth = [
            transaction
            for transaction in aggregation.bank_transactions
            if transaction.summary.strip() == "理财"
        ]
        financial_payments = [
            transaction
            for transaction in aggregation.bank_transactions
            if transaction.summary.strip() == "金融付款"
        ]
        wealth_expense = -sum(transaction.amount_minor for transaction in wealth)
        financial_payment_expense = -sum(
            transaction.amount_minor for transaction in financial_payments
        )
        note = "起点早于 PDF，区间不完整" if aggregation.partial_start else ""
        lines.append(
            f"| {aggregation.online.occurred_at:%Y-%m-%d %H:%M:%S} | "
            f"{len(wealth)} / {format_money(wealth_expense)} | "
            f"{len(financial_payments)} / {format_money(financial_payment_expense)} | "
            f"{format_money(aggregation.bank_expense_minor)} | "
            f"{format_money(aggregation.online_expense_minor)} | "
            f"{format_money(aggregation.expense_delta_minor)} | {note} |"
        )

    lines.extend(
        [
            "",
            f"- 已分配 PDF 理财类支出：{format_money(allocated_finance_expense)} CNY",
            f"- 线上汇总转账：{format_money(online_finance_expense)} CNY",
            f"- 可完整比较的 {len(complete_finance)} 个区间：PDF {format_money(complete_finance_bank_expense)} CNY，线上 {format_money(complete_online_finance_expense)} CNY，差额 {format_money(complete_finance_bank_expense - complete_online_finance_expense)} CNY",
            (
                f"- 首个不完整区间差额：{format_money(partial_finance_delta)} CNY（仅表示 PDF 起始日前数据未覆盖）"
                if partial_finance
                else "- 首个区间：已按日拆分，PDF 覆盖完整"
            ),
            f"- 末次汇总后待汇总：{len(finance.pending_bank)} 笔，{format_money(pending_finance_expense)} CNY",
            f"- 已排除账单期前理财残差：{len(finance.excluded_residual_online)} 笔，{format_money(residual_finance_expense)} CNY",
            f"- 整个 PDF 覆盖范围（含待汇总）：PDF {format_money(full_scope_finance_expense)} CNY，线上 {format_money(online_finance_expense)} CNY，差额 {format_money(full_scope_finance_expense - online_finance_expense)} CNY",
            "",
            "## 同日金额差异（完整）",
            "",
            f"双方交易时间位于同一自然日；普通交易允许 {format_money(amount_tolerance_minor)} CNY 差异，双方均不超过 {format_money(small_amount_limit_minor)} CNY 时不限制差额。",
            "",
            "| 银行时间 | 银行金额 | 线上时间 | 线上金额 | 时间差 | 金额差（银行 - 线上） | 交易 ID |",
            "|---|---:|---|---:|---:|---:|---|",
        ]
    )
    for match in amount_difference_matches:
        lines.append(
            f"| {match.bank.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(match.bank.amount_minor)} | "
            f"{match.online.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(match.online.amount_minor)} | "
            f"{match.time_delta_seconds} 秒 | {format_money(match.amount_delta_minor)} | "
            f"`{match.online.transaction_id}` |"
        )
    if not amount_difference_matches:
        lines.append("| — | — | — | — | — | — | — |")

    lines.extend(
        [
            "",
            "## 原始按日差异",
            "",
            "仅列出笔数、收入或支出任一项不一致的日期；理财汇总转账的记账日期未回拨到其所覆盖的 PDF 日期。",
            "",
            "| 日期 | 银行笔数 | 线上笔数 | 收入差额 | 支出差额 | 净额差额 |",
            "|---|---:|---:|---:|---:|---:|",
        ]
    )

    for day in sorted(set(bank_daily) | set(online_daily)):
        bank_day = bank_daily.get(day, (0, 0, 0))
        online_day = online_daily.get(day, (0, 0, 0))
        if bank_day == online_day:
            continue
        lines.append(
            f"| {day} | {bank_day[0]} | {online_day[0]} | "
            f"{format_money(bank_day[1] - online_day[1])} | "
            f"{format_money(bank_day[2] - online_day[2])} | "
            f"{format_money((bank_day[1] - bank_day[2]) - (online_day[1] - online_day[2]))} |"
        )

    lines.extend(["", "## 其他容差匹配与未匹配明细（节选）", ""])
    lines.extend(
        [
            f"### 已排除的等额转入/转出（完整，共 {len(suppressed_bank_pairs)} 组）",
            "",
            "| 转出时间 | 转出金额 | 转入时间 | 转入金额 | 抵销金额 |",
            "|---|---:|---|---:|---:|",
        ]
    )
    for pair in suppressed_bank_pairs:
        lines.append(
            f"| {pair.outflow.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(pair.outflow.amount_minor)} | "
            f"{pair.inflow.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(pair.inflow.amount_minor)} | "
            f"{format_money(abs(pair.outflow.amount_minor))} |"
        )
    if not suppressed_bank_pairs:
        lines.append("| — | — | — | — | — |")
    lines.append("")
    if detail_limit <= 0:
        lines.append("命令参数要求省略逐笔明细。")
    else:
        lines.extend(
            [
                f"### 金额相同、仅时间偏差（最多 {detail_limit} 笔）",
                "",
                "| 银行时间 | 银行金额 | 线上时间 | 线上金额 | 时间差 | 金额差（银行 - 线上） | 交易 ID |",
                "|---|---:|---|---:|---:|---:|---|",
            ]
        )
        for match in [
            item for item in result.near_matches if item.amount_delta_minor == 0
        ][:detail_limit]:
            lines.append(
                f"| {match.bank.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(match.bank.amount_minor)} | "
                f"{match.online.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(match.online.amount_minor)} | "
                f"{match.time_delta_seconds} 秒 | {format_money(match.amount_delta_minor)} | "
                f"`{match.online.transaction_id}` |"
            )
        if not near_same_amount:
            lines.append("| — | — | — | — | — | — | — |")

        lines.extend(
            [
                "",
                f"### 银行有、线上无（已排除理财汇总，最多 {detail_limit} 笔）",
                "",
                "| 时间 | 金额 | PDF 位置 |",
                "|---|---:|---|",
            ]
        )
        for transaction in result.bank_only[:detail_limit]:
            lines.append(
                f"| {transaction.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(transaction.amount_minor)} | 第 {transaction.page} 页第 {transaction.row} 行 |"
            )
        if not result.bank_only:
            lines.append("| — | — | — |")

        lines.extend(
            [
                "",
                f"### 线上有、银行无（已排除理财汇总，最多 {detail_limit} 笔）",
                "",
                "| 时间 | 金额 | 类型 | 备注 | 交易 ID |",
                "|---|---:|---:|---|---|",
            ]
        )
        for transaction in result.online_only[:detail_limit]:
            comment = transaction.comment.replace("|", "\\|").replace("\n", " ")
            lines.append(
                f"| {transaction.occurred_at:%Y-%m-%d %H:%M:%S} | {format_money(transaction.amount_minor)} | "
                f"{transaction.transaction_type} | {comment} | `{transaction.transaction_id}` |"
            )
        if not result.online_only:
            lines.append("| — | — | — | — | — |")

    lines.extend(
        [
            "",
            "## 核对规则与限制",
            "",
            "- PDF 先过滤 7pt 正文字符，再按表格线提取日期时间、收支金额和余额；逐笔验证余额连续性。",
            f"- PDF 摘要为“理财”或“金融付款”的支出按线上转入“{fund_account_name}”的时间分段汇总；首个区间和末次转账后的待汇总区间单独标记。",
            f"- 备注以“{DAILY_FINANCE_PREFIX}”开头的转账按其具体日期核对；备注为“{PRE_STATEMENT_RESIDUAL_COMMENT}”的转账属于 PDF 起始日前金额，从本期统计中排除。",
            f"- 线上接口按整天查询；普通流水只保留 PDF 精确起止时间前后 {time_tolerance_seconds} 秒内的数据，理财汇总转账仍按其实际汇总时间参与区间核对。",
            "- 先按“UTC+8 本地交易时间（精确到秒）+ 有符号分金额”做精确多重集匹配，再对剩余记录做一对一容差匹配。",
            f"- 容差匹配要求收支方向相同且位于同一自然日，候选先按时间差、再按金额差排序；普通金额容差为 {format_money(amount_tolerance_minor)} CNY，双方均不超过 {format_money(small_amount_limit_minor)} CNY 时允许金额不同并完整披露差额。",
            "- 银行侧未匹配流水中，金额绝对值相同且符号相反的转入/转出按时间最接近原则一对一抵销，并从未匹配统计中排除。",
            "- 线上收入记正、支出记负；转账按目标账户是转出方或转入方确定符号。",
            "- 未传 `--account` 时，仅在 PDF 期末余额与某个线上 CNY 账户余额唯一一致时自动选择；否则拒绝猜测。",
            "- 该脚本不会调用交易新增、修改、删除、导入或账户修改接口。",
            "",
        ]
    )
    return "\n".join(lines)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "statement", type=Path, help="ICBC statement PDF or converted XLSX"
    )
    parser.add_argument(
        "--base-url", default="https://example.com", help="ezBookkeeping site URL"
    )
    parser.add_argument(
        "--username", help="ezBookkeeping login name; omit for PDF-only validation"
    )
    parser.add_argument("--account", help="exact ezBookkeeping account name or ID")
    parser.add_argument("--output", type=Path, help="Markdown report path")
    parser.add_argument(
        "--detail-limit", type=int, default=50, help="maximum unmatched rows per side"
    )
    parser.add_argument(
        "--time-tolerance",
        type=int,
        default=10,
        help="statement-boundary tolerance in seconds; near matches use the same day",
    )
    parser.add_argument(
        "--amount-tolerance-cents",
        type=int,
        default=5,
        help="near-match tolerance in cents",
    )
    parser.add_argument(
        "--small-amount-limit-cents",
        type=int,
        default=1000,
        help="match by time when both amounts do not exceed this many cents",
    )
    parser.add_argument(
        "--fund-account",
        default="支付宝基金",
        help="exact destination account name for aggregated finance transfers",
    )
    parser.add_argument(
        "--strict", action="store_true", help="exit 2 when unmatched rows exist"
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    statement_path = args.statement.resolve()
    if not statement_path.is_file():
        raise ReconcileError(f"statement does not exist: {statement_path}")

    bank_transactions = load_statement(statement_path)
    if not args.username:
        count, income, expense = summarize(bank_transactions)
        print(
            f"Statement validated: {count} transactions, income {format_money(income)}, "
            f"expense {format_money(expense)}, closing balance "
            f"{format_money(bank_transactions[-1].balance_minor)} CNY"
        )
        return 0

    password = os.environ.get("EZBOOKKEEPING_PASSWORD") or getpass.getpass(
        "ezBookkeeping password: "
    )
    client = EzBookkeepingClient(args.base_url)
    client.authorize(args.username, password)
    accounts = client.list_accounts()
    account = select_account(
        accounts, args.account, bank_transactions[-1].balance_minor
    )

    first_day = datetime.combine(
        bank_transactions[0].occurred_at.date(), time.min, LOCAL_ZONE
    )
    last_day = datetime.combine(
        bank_transactions[-1].occurred_at.date(), time.max, LOCAL_ZONE
    )
    rows = client.list_transactions(
        str(account["id"]), int(first_day.timestamp()), int(last_day.timestamp())
    )
    online_transactions, unsupported = normalize_online_transactions(
        rows, str(account["id"])
    )
    fund_accounts = [
        item
        for item in accounts
        if str(item.get("name", "")).strip() == args.fund_account.strip()
    ]
    if len(fund_accounts) != 1:
        raise ReconcileError(
            f"expected one fund account named {args.fund_account!r}, found {len(fund_accounts)}"
        )
    finance = reconcile_finance_aggregations(
        bank_transactions,
        online_transactions,
        str(account["id"]),
        {str(fund_accounts[0]["id"])},
    )
    regular_online, outside_period_online = partition_online_by_period(
        finance.regular_online,
        bank_transactions[0].occurred_at,
        bank_transactions[-1].occurred_at,
        args.time_tolerance,
    )
    excluded_online_ids = {
        id(transaction)
        for transaction in outside_period_online + finance.excluded_residual_online
    }
    online_transactions = [
        transaction
        for transaction in online_transactions
        if id(transaction) not in excluded_online_ids
    ]
    lower_timestamp = int(
        (
            bank_transactions[0].occurred_at - timedelta(seconds=args.time_tolerance)
        ).timestamp()
    )
    upper_timestamp = int(
        (
            bank_transactions[-1].occurred_at + timedelta(seconds=args.time_tolerance)
        ).timestamp()
    )
    unsupported = [
        row
        for row in unsupported
        if lower_timestamp <= int(row.get("time", 0)) <= upper_timestamp
    ]
    result = reconcile(
        finance.regular_bank,
        regular_online,
        args.amount_tolerance_cents,
        args.small_amount_limit_cents,
    )
    remaining_bank_only, suppressed_bank_pairs = suppress_equal_opposite_bank_only(
        result.bank_only
    )
    result = Reconciliation(
        result.exact_matched_count,
        result.near_matches,
        remaining_bank_only,
        result.online_only,
    )
    report = render_report(
        bank_transactions,
        online_transactions,
        unsupported,
        finance,
        result,
        suppressed_bank_pairs,
        account,
        args.base_url,
        args.detail_limit,
        args.time_tolerance,
        args.amount_tolerance_cents,
        args.small_amount_limit_cents,
        args.account,
        args.fund_account,
        len(outside_period_online),
    )

    output_path = args.output or statement_path.with_name(
        f"{statement_path.stem}_核对报告.md"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(report, encoding="utf-8", newline="\n")
    print(
        f"Report written to {output_path.resolve()} | exact={result.exact_matched_count} "
        f"near={len(result.near_matches)} "
        f"finance_periods={len(finance.aggregations)} "
        f"outside_period={len(outside_period_online)} "
        f"bank_only={len(result.bank_only)} online_only={len(result.online_only)}"
    )
    return (
        2
        if args.strict
        and (
            result.bank_only
            or result.online_only
            or unsupported
            or finance.pending_bank
            or any(item.expense_delta_minor for item in finance.aggregations)
        )
        else 0
    )


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ReconcileError as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1)
