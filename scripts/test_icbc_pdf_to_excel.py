import importlib.util
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, main

from openpyxl import load_workbook

SCRIPTS = Path(__file__).parent
sys.path.insert(0, str(SCRIPTS))
SCRIPT = SCRIPTS / "icbc_pdf_to_excel.py"
SPEC = importlib.util.spec_from_file_location("icbc_pdf_to_excel", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


class IcbcPdfToExcelTest(TestCase):
    def setUp(self):
        self.when = datetime(
            2026, 3, 16, 9, 51, 37, tzinfo=timezone(timedelta(hours=8), "CST")
        )
        self.transactions = [
            MODULE.BankTransaction(
                self.when,
                -240,
                100000,
                1,
                1,
                "消费",
                "6222",
                "活期",
                "serial-1",
                "人民币",
                "钞",
                "北京",
                "快捷支付",
            ),
            MODULE.BankTransaction(
                self.when + timedelta(seconds=1),
                1000,
                101000,
                1,
                2,
                "利息",
                "6222",
                "活期",
                "serial-2",
                "人民币",
                "钞",
                "北京",
                "柜面",
            ),
        ]

    def test_workbook_round_trip_preserves_statement(self):
        with TemporaryDirectory() as directory:
            pdf = Path(directory) / "source.pdf"
            pdf.write_bytes(b"test-source")
            output = Path(directory) / "statement.xlsx"
            MODULE.write_statement_excel(self.transactions, pdf, output)
            self.assertEqual(MODULE.parse_statement_excel(output), self.transactions)

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["转换说明", "工商银行流水"])
            formulas = [
                cell.value
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            self.assertEqual(formulas, [])
            self.assertEqual(workbook["转换说明"]["B3"].value, "source.pdf")
            self.assertEqual(
                workbook["工商银行流水"]["J2"].number_format, MODULE.MONEY_FORMAT
            )
            workbook.close()


if __name__ == "__main__":
    main()
